"""Extract plain text + whatever metadata hints we can get out of a
downloaded file, per file type. Never raises for a single bad file — the
indexer treats that file as failed and moves on.
"""

from __future__ import annotations

import sys
from dataclasses import dataclass
from html.parser import HTMLParser
from pathlib import Path
from typing import Callable, Dict, Optional

MAX_CHARS = 8000
MAX_PDF_PAGES = 4
MAX_DOCX_PARAGRAPHS = 300


@dataclass
class ExtractedDocument:
    text: str
    title_hint: Optional[str]
    author_hint: Optional[str]
    file_type: str


def _clean_meta(value: Optional[str]) -> Optional[str]:
    if not value:
        return None
    value = str(value).strip()
    return value or None


def extract_pdf(path: Path) -> ExtractedDocument:
    try:
        import pypdf
    except ImportError as exc:
        raise RuntimeError("Reading PDFs requires `pip install pypdf`") from exc

    reader = pypdf.PdfReader(str(path))
    chunks = []
    for page in reader.pages[:MAX_PDF_PAGES]:
        try:
            chunks.append(page.extract_text() or "")
        except Exception:
            continue
    text = "\n".join(chunks)[:MAX_CHARS]

    meta = reader.metadata
    title_hint = _clean_meta(getattr(meta, "title", None)) if meta else None
    author_hint = _clean_meta(getattr(meta, "author", None)) if meta else None
    return ExtractedDocument(text=text, title_hint=title_hint, author_hint=author_hint, file_type="pdf")


def extract_docx(path: Path) -> ExtractedDocument:
    try:
        import docx
    except ImportError as exc:
        raise RuntimeError("Reading .docx files requires `pip install python-docx`") from exc

    document = docx.Document(str(path))
    paragraphs = [p.text for p in document.paragraphs[:MAX_DOCX_PARAGRAPHS] if p.text.strip()]
    text = "\n".join(paragraphs)[:MAX_CHARS]

    core = document.core_properties
    title_hint = _clean_meta(core.title)
    author_hint = _clean_meta(core.author)
    return ExtractedDocument(text=text, title_hint=title_hint, author_hint=author_hint, file_type="docx")


def extract_txt(path: Path) -> ExtractedDocument:
    text = path.read_text(encoding="utf-8", errors="ignore")[:MAX_CHARS]
    return ExtractedDocument(text=text, title_hint=None, author_hint=None, file_type="txt")


class _HtmlTextExtractor(HTMLParser):
    def __init__(self):
        super().__init__()
        self._chunks = []
        self._title_chunks = []
        self._in_title = False

    def handle_starttag(self, tag, attrs):
        if tag == "title":
            self._in_title = True

    def handle_endtag(self, tag):
        if tag == "title":
            self._in_title = False

    def handle_data(self, data):
        (self._title_chunks if self._in_title else self._chunks).append(data)

    @property
    def text(self) -> str:
        return " ".join(" ".join(self._chunks).split())

    @property
    def title(self) -> Optional[str]:
        title = " ".join(" ".join(self._title_chunks).split())
        return title or None


def extract_html(path: Path) -> ExtractedDocument:
    raw = path.read_text(encoding="utf-8", errors="ignore")
    parser = _HtmlTextExtractor()
    parser.feed(raw)
    return ExtractedDocument(
        text=parser.text[:MAX_CHARS],
        title_hint=parser.title,
        author_hint=None,
        file_type="html",
    )


_EXTRACTORS: Dict[str, Callable[[Path], ExtractedDocument]] = {
    ".pdf": extract_pdf,
    ".docx": extract_docx,
    ".txt": extract_txt,
    ".html": extract_html,
    ".htm": extract_html,
}


def extract_document(path: Path) -> Optional[ExtractedDocument]:
    """Dispatch to the right extractor for `path`'s extension.

    Returns None (rather than raising) if the extension is unsupported or
    extraction fails — the caller logs and skips.
    """
    extractor = _EXTRACTORS.get(path.suffix.lower())
    if extractor is None:
        return None
    try:
        return extractor(path)
    except Exception as exc:
        print(f"  [extract] failed to read {path} ({exc})", file=sys.stderr)
        return None
