"""The criticality chart and reading log — the two working documents the
dissertation bootcamp material recommends keeping while reviewing."""

from types import SimpleNamespace
from unittest.mock import patch

from openpyxl import Workbook

from thesis_tools.library.index_store import LibraryEntry, LibraryIndex
from thesis_tools.library.literature_matrix import (
    CRITICAL_QUESTIONS,
    MAX_CRITICALITY_PAPERS,
    READING_LOG_COLUMNS,
    build_literature_matrix_workbook,
)
from thesis_tools.sources.base import Paper

QUESTION = "How do travellers use AI tools?"

_SPECS = [
    ("Traveller trust in generative AI tools", 2024, "Generative AI shapes travel planning trust.", "supports"),
    ("Generative AI and travel planning", 2025, "Consumer trust in travel planning with AI.", "supports"),
    ("No measurable OTA share loss", 2015, "No measurable change in OTA share from AI.", "challenges"),
    ("Revenue management with machine learning", 2024, "Machine learning for hotel revenue management.", "unrelated"),
]


class _FakeAnalysis:
    def __init__(self, specs=_SPECS):
        self._by_title = {t: st for t, _, _, st in specs}
        self._specs = specs

    def stances_for(self, paper):
        stance = self._by_title.get(paper.title, "unrelated")
        return {} if stance == "unrelated" else {QUESTION: SimpleNamespace(stance=stance, reason="x")}

    def titles_by_stance(self, question):
        grouped = {"supports": [], "challenges": [], "mixed": []}
        for title, _, _, stance in self._specs:
            if stance in grouped:
                grouped[stance].append(title)
        return grouped


def _index(specs=_SPECS):
    return LibraryIndex([
        LibraryEntry(
            file_path=f"/lib/p{i}.pdf", file_hash=str(i), file_type="pdf", size_bytes=1,
            indexed_at="2026-01-01T00:00:00", confidence="verified-doi", doi=f"10.1/{i}",
            paper=Paper(title=t, year=y, authors=[f"Ada Lovelace{i}"], doi=f"10.1/{i}", abstract=a),
        )
        for i, (t, y, a, _) in enumerate(specs)
    ])


def _workbook(specs=_SPECS, themes=None):
    return build_literature_matrix_workbook(
        _index(specs),
        sub_questions=[QUESTION],
        research_question="How does AI affect travel planning?",
        analysis=_FakeAnalysis(specs),
        themes=themes,
    )


def _column(ws, header_row_value):
    """The question row, as a list of its cells' values."""
    for row in range(1, ws.max_row + 1):
        if ws.cell(row=row, column=1).value == header_row_value:
            return [ws.cell(row=row, column=c).value for c in range(2, ws.max_column + 1)]
    raise AssertionError(f"no row for {header_row_value!r}")


def test_the_workbook_carries_all_three_sheets():
    assert _workbook().sheetnames == ["Literature Review Matrix", "Criticality", "Reading log"]


def test_the_criticality_chart_asks_every_question_of_every_paper():
    ws = _workbook()["Criticality"]
    asked = [ws.cell(row=r, column=1).value for r in range(2, ws.max_row + 1)]
    assert asked == CRITICAL_QUESTIONS


def test_questions_only_the_student_can_answer_are_left_blank():
    """A blank cell is the question being put to the reader. Pre-filling "is
    there bias?" with a guess would replace judgement with the appearance of
    it."""
    ws = _workbook()["Criticality"]
    for question in ("Is the research biased?", "Is the evidence robust?", "What is the main argument?"):
        assert all(not value for value in _column(ws, question)), question


def test_what_the_tool_actually_knows_is_filled_in():
    ws = _workbook()["Criticality"]

    relevance = _column(ws, "Is the research relevant?")
    assert all("overlap with the research question" in (v or "") for v in relevance)
    assert all("your call" in (v or "") for v in relevance)  # a prompt, not a verdict

    dated = _column(ws, "Is the research or theory a product of its time?")
    assert any("2015" in (v or "") and "year(s) old" in (v or "") for v in dated)


def test_the_agreement_cell_names_who_rather_than_restating_the_stance():
    """The one critical question the toolkit answers better than memory: it
    has already classified every paper against every sub-question."""
    ws = _workbook()["Criticality"]
    cells = [v for v in _column(ws, "Who agrees or disagrees with this stance?") if v]

    assert cells
    supporter = next(v for v in cells if "Supports" in v)
    assert "agrees with:" in supporter
    assert "disagrees with: No measurable OTA share loss" in supporter


def test_a_paper_nobody_else_takes_a_position_on_says_so():
    specs = [("Alone on this", 2024, "A solitary claim about travel planning.", "supports")]
    ws = _workbook(specs)["Criticality"]
    cell = next(v for v in _column(ws, "Who agrees or disagrees with this stance?") if v)
    assert "no other paper in the library takes a position" in cell


def test_themes_are_carried_across_from_the_visualization():
    themes = [{"label": "travel planning", "papers": [{"key": _index().entries[0].paper.key()}]}]
    ws = _workbook(themes=themes)["Criticality"]
    assert any("travel planning" in (v or "") for v in _column(ws, "What are the main themes?"))


def test_the_chart_is_capped_at_a_workable_number_of_papers():
    """Eighty columns is not a working document."""
    many = [(f"Paper {i} on travel planning", 2024, "Travel planning.", "supports") for i in range(40)]
    ws = _workbook(many)["Criticality"]
    assert ws.max_column - 1 == MAX_CRITICALITY_PAPERS


def test_the_reading_log_marks_which_columns_are_the_students_to_write():
    ws = _workbook()["Reading log"]
    headers = [ws.cell(row=1, column=c).value for c in range(1, ws.max_column + 1)]

    assert "Text (exact) (yours to write)" in headers
    assert "What it means to me (yours to write)" in headers
    assert "Title" in headers  # pre-filled ones are not labelled


def test_the_reading_log_prefills_only_what_is_known():
    ws = _workbook()["Reading log"]
    by_header = {
        ws.cell(row=1, column=c).value: ws.cell(row=2, column=c).value
        for c in range(1, ws.max_column + 1)
    }

    assert by_header["Title"]
    assert "2024" in by_header["Author / Date"]
    assert by_header["How it relates to what others say"]  # seeded from the stance analysis
    # The reflection is the point of keeping the log; writing it here would
    # defeat the exercise.
    assert not by_header["Text (exact) (yours to write)"]
    assert not by_header["What it means to me (yours to write)"]


def test_every_indexed_paper_gets_a_reading_log_row():
    ws = _workbook()["Reading log"]
    assert ws.max_row == len(_SPECS) + 1


def test_the_sheets_survive_a_library_with_no_questions_or_analysis():
    wb = build_literature_matrix_workbook(_index())
    assert "Criticality" in wb.sheetnames and "Reading log" in wb.sheetnames


def test_an_empty_library_does_not_get_an_empty_criticality_chart():
    wb = build_literature_matrix_workbook(LibraryIndex([]))
    assert "Criticality" not in wb.sheetnames
