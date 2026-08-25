from unittest.mock import patch

from thesis_tools.cli import main
from thesis_tools.library.index_store import LibraryEntry, LibraryIndex
from thesis_tools.sources.base import Paper


def test_visualize_library_errors_when_no_index(tmp_path, capsys):
    missing = tmp_path / "library" / "index.json"
    rc = main(["visualize-library", "--index-path", str(missing)])
    assert rc == 2
    assert "no index found" in capsys.readouterr().err


def test_visualize_library_writes_html(tmp_path):
    index_path = tmp_path / "library" / "index.json"
    output_path = tmp_path / "library" / "visualization.html"

    index = LibraryIndex(
        [
            LibraryEntry(
                file_path="a.pdf",
                file_hash="hash-a",
                file_type="pdf",
                size_bytes=100,
                indexed_at="2026-01-01T00:00:00",
                confidence="verified-doi",
                doi="10.1/a",
                paper=Paper(title="A Paper", doi="10.1/a", year=2023, sources=["crossref"]),
            )
        ]
    )
    index.save(index_path)

    rc = main(["visualize-library", "--index-path", str(index_path), "--output", str(output_path)])

    assert rc == 0
    assert output_path.is_file()
    html = output_path.read_text(encoding="utf-8")
    assert "Library Visualization" in html
    assert "Crossref" in html


def test_visualize_library_default_output_path(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    index_path = tmp_path / "library" / "index.json"
    LibraryIndex().save(index_path)

    rc = main(["visualize-library", "--index-path", str(index_path)])

    assert rc == 0
    assert (tmp_path / "library" / "visualization.html").is_file()


def _index_with_paper(index_path, title="Relevant Paper", full_text_excerpt=None):
    entry = LibraryEntry(
        file_path="a.pdf",
        file_hash="hash-a",
        file_type="pdf",
        size_bytes=100,
        indexed_at="2026-01-01T00:00:00",
        confidence="verified-doi",
        doi="10.1/a",
        paper=Paper(title=title, doi="10.1/a", year=2023, sources=["crossref"], full_text_excerpt=full_text_excerpt),
    )
    LibraryIndex([entry]).save(index_path)


def test_visualize_library_anchors_on_explicit_subquestions_flag(tmp_path):
    index_path = tmp_path / "library" / "index.json"
    output_path = tmp_path / "library" / "visualization.html"
    project_file = tmp_path / "project.json"
    _index_with_paper(
        index_path,
        title="Relevant Paper",
        full_text_excerpt="This significantly supports the claim that X affects Y, consistent with theory.",
    )

    rc = main(
        [
            "visualize-library",
            "--index-path", str(index_path),
            "--output", str(output_path),
            "--project-file", str(project_file),
            "--sub-questions", "Does X affect Y?",
        ]
    )

    assert rc == 0
    html = output_path.read_text(encoding="utf-8")
    assert "Does X affect Y?" in html
    assert "Relevant Paper" in html


def test_visualize_library_reuses_project_subquestions_when_flag_omitted(tmp_path):
    from thesis_tools.project import ProjectState

    index_path = tmp_path / "library" / "index.json"
    output_path = tmp_path / "library" / "visualization.html"
    project_file = tmp_path / "project.json"
    ProjectState(sub_questions=["Saved question from Part 1?"]).save(str(project_file))
    _index_with_paper(index_path)

    rc = main(
        [
            "visualize-library",
            "--index-path", str(index_path),
            "--output", str(output_path),
            "--project-file", str(project_file),
        ]
    )

    assert rc == 0
    html = output_path.read_text(encoding="utf-8")
    assert "Saved question from Part 1?" in html


def test_visualize_library_no_subquestions_shows_prompt(tmp_path):
    index_path = tmp_path / "library" / "index.json"
    output_path = tmp_path / "library" / "visualization.html"
    project_file = tmp_path / "project.json"
    _index_with_paper(index_path)

    rc = main(
        [
            "visualize-library",
            "--index-path", str(index_path),
            "--output", str(output_path),
            "--project-file", str(project_file),
        ]
    )

    assert rc == 0
    html = output_path.read_text(encoding="utf-8")
    assert "No sub-questions configured for this run" in html


def test_visualize_library_flags_low_relevance_papers_with_explicit_question(tmp_path):
    index_path = tmp_path / "library" / "index.json"
    output_path = tmp_path / "library" / "visualization.html"
    project_file = tmp_path / "project.json"
    _index_with_paper(index_path, title="A Completely Unrelated Coffee Farming Study")

    rc = main(
        [
            "visualize-library",
            "--index-path", str(index_path),
            "--output", str(output_path),
            "--project-file", str(project_file),
            "--question", "digital transformation sustainability targets",
        ]
    )

    assert rc == 0
    html = output_path.read_text(encoding="utf-8")
    assert "low relevance to your research question" in html


@patch("thesis_tools.subquestions.llm.get_client")
@patch("thesis_tools.subquestions.llm.ask")
def test_visualize_library_caches_stance_classification_by_default(mock_ask, mock_get_client, tmp_path):
    mock_get_client.return_value = object()
    mock_ask.return_value = "1: supports - confirms it."
    index_path = tmp_path / "library" / "index.json"
    output_path = tmp_path / "library" / "visualization.html"
    project_file = tmp_path / "project.json"
    _index_with_paper(
        index_path,
        title="Relevant Paper",
        full_text_excerpt="This significantly supports the claim that X affects Y, consistent with theory.",
    )
    common_args = [
        "visualize-library",
        "--index-path", str(index_path),
        "--output", str(output_path),
        "--project-file", str(project_file),
        "--sub-questions", "Does X affect Y?",
        "--llm-summaries",
    ]

    assert main(common_args) == 0
    assert mock_ask.call_count == 1
    # Cache file defaults to a sibling of --index-path.
    assert (tmp_path / "library" / "stance_cache.json").is_file()

    assert main(common_args) == 0
    assert mock_ask.call_count == 1  # unchanged paper + unchanged question -> no second call


@patch("thesis_tools.subquestions.llm.get_client")
@patch("thesis_tools.subquestions.llm.ask")
def test_visualize_library_no_stance_cache_always_reclassifies(mock_ask, mock_get_client, tmp_path):
    mock_get_client.return_value = object()
    mock_ask.return_value = "1: supports - confirms it."
    index_path = tmp_path / "library" / "index.json"
    output_path = tmp_path / "library" / "visualization.html"
    project_file = tmp_path / "project.json"
    _index_with_paper(
        index_path,
        title="Relevant Paper",
        full_text_excerpt="This significantly supports the claim that X affects Y, consistent with theory.",
    )
    common_args = [
        "visualize-library",
        "--index-path", str(index_path),
        "--output", str(output_path),
        "--project-file", str(project_file),
        "--sub-questions", "Does X affect Y?",
        "--llm-summaries",
        "--no-stance-cache",
    ]

    assert main(common_args) == 0
    assert mock_ask.call_count == 1
    assert not (tmp_path / "library" / "stance_cache.json").exists()

    assert main(common_args) == 0
    assert mock_ask.call_count == 2  # caching disabled -> reclassified every run


@patch("thesis_tools.subquestions.llm.get_client")
@patch("thesis_tools.subquestions.llm.ask")
def test_visualize_library_explicit_stance_cache_path(mock_ask, mock_get_client, tmp_path):
    mock_get_client.return_value = object()
    mock_ask.return_value = "1: supports - confirms it."
    index_path = tmp_path / "library" / "index.json"
    output_path = tmp_path / "library" / "visualization.html"
    project_file = tmp_path / "project.json"
    custom_cache = tmp_path / "elsewhere" / "cache.json"
    _index_with_paper(
        index_path,
        title="Relevant Paper",
        full_text_excerpt="This significantly supports the claim that X affects Y, consistent with theory.",
    )

    rc = main(
        [
            "visualize-library",
            "--index-path", str(index_path),
            "--output", str(output_path),
            "--project-file", str(project_file),
            "--sub-questions", "Does X affect Y?",
            "--llm-summaries",
            "--stance-cache-path", str(custom_cache),
        ]
    )

    assert rc == 0
    assert custom_cache.is_file()
    assert not (tmp_path / "library" / "stance_cache.json").exists()


def test_visualize_library_writes_matrix_next_to_output_by_default(tmp_path):
    index_path = tmp_path / "library" / "index.json"
    output_path = tmp_path / "library" / "visualization.html"
    _index_with_paper(index_path, title="A Paper")

    rc = main(["visualize-library", "--index-path", str(index_path), "--output", str(output_path)])

    assert rc == 0
    matrix_path = tmp_path / "library" / "literature_matrix.xlsx"
    assert matrix_path.is_file()

    from openpyxl import load_workbook

    ws = load_workbook(matrix_path).active
    header, row = list(ws.iter_rows(values_only=True))
    assert row[4] == "A Paper"

    html = output_path.read_text(encoding="utf-8")
    assert 'href="literature_matrix.xlsx" download' in html


def test_visualize_library_custom_matrix_path(tmp_path):
    index_path = tmp_path / "library" / "index.json"
    output_path = tmp_path / "library" / "visualization.html"
    custom_matrix = tmp_path / "exports" / "matrix.xlsx"
    _index_with_paper(index_path, title="A Paper")

    rc = main(
        [
            "visualize-library",
            "--index-path", str(index_path),
            "--output", str(output_path),
            "--matrix-path", str(custom_matrix),
        ]
    )

    assert rc == 0
    assert custom_matrix.is_file()
    assert not (tmp_path / "library" / "literature_matrix.xlsx").exists()

    html = output_path.read_text(encoding="utf-8")
    # Relative from library/ to exports/matrix.xlsx.
    assert 'href="../exports/matrix.xlsx" download' in html


def test_visualize_library_no_matrix_skips_excel_and_link(tmp_path):
    index_path = tmp_path / "library" / "index.json"
    output_path = tmp_path / "library" / "visualization.html"
    _index_with_paper(index_path, title="A Paper")

    rc = main(
        [
            "visualize-library",
            "--index-path", str(index_path),
            "--output", str(output_path),
            "--no-matrix",
        ]
    )

    assert rc == 0
    assert not (tmp_path / "library" / "literature_matrix.xlsx").exists()
    html = output_path.read_text(encoding="utf-8")
    assert "Download as Excel" not in html


def test_visualize_library_matrix_uses_configured_style(tmp_path):
    index_path = tmp_path / "library" / "index.json"
    output_path = tmp_path / "library" / "visualization.html"
    entry = LibraryEntry(
        file_path="a.pdf",
        file_hash="hash-a",
        file_type="pdf",
        size_bytes=100,
        indexed_at="2026-01-01T00:00:00",
        confidence="verified-doi",
        doi="10.1/a",
        paper=Paper(title="A Paper", authors=["Jane Doe"], doi="10.1/a", year=2023, sources=["crossref"]),
    )
    LibraryIndex([entry]).save(index_path)

    rc = main(
        [
            "visualize-library",
            "--index-path", str(index_path),
            "--output", str(output_path),
            "--style", "ieee",
        ]
    )

    assert rc == 0
    from openpyxl import load_workbook

    ws = load_workbook(tmp_path / "library" / "literature_matrix.xlsx").active
    _, row = list(ws.iter_rows(values_only=True))
    assert row[1].startswith("[1]")


def test_visualize_library_empty_index_writes_no_matrix(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    index_path = tmp_path / "library" / "index.json"
    LibraryIndex().save(index_path)

    rc = main(["visualize-library", "--index-path", str(index_path)])

    assert rc == 0
    assert not (tmp_path / "library" / "literature_matrix.xlsx").exists()


@patch("thesis_tools.subquestions.llm.get_client")
@patch("thesis_tools.subquestions.llm.ask")
def test_visualize_library_matrix_does_not_trigger_extra_llm_calls(mock_ask, mock_get_client, tmp_path):
    mock_get_client.return_value = object()
    mock_ask.return_value = "1: supports - confirms it."
    index_path = tmp_path / "library" / "index.json"
    output_path = tmp_path / "library" / "visualization.html"
    project_file = tmp_path / "project.json"
    _index_with_paper(
        index_path,
        title="Relevant Paper",
        full_text_excerpt="This significantly supports the claim that X affects Y, consistent with theory.",
    )

    rc = main(
        [
            "visualize-library",
            "--index-path", str(index_path),
            "--output", str(output_path),
            "--project-file", str(project_file),
            "--sub-questions", "Does X affect Y?",
            "--llm-summaries",
        ]
    )

    assert rc == 0
    # One classification call for the HTML stats — building the matrix from
    # the same stats must not trigger a second round of Claude calls.
    assert mock_ask.call_count == 1
    assert (tmp_path / "library" / "literature_matrix.xlsx").is_file()


def test_visualize_library_reuses_project_research_question(tmp_path):
    from thesis_tools.project import ProjectState

    index_path = tmp_path / "library" / "index.json"
    output_path = tmp_path / "library" / "visualization.html"
    project_file = tmp_path / "project.json"
    ProjectState(research_question="digital transformation sustainability targets").save(str(project_file))
    _index_with_paper(index_path, title="A Completely Unrelated Coffee Farming Study")

    rc = main(
        [
            "visualize-library",
            "--index-path", str(index_path),
            "--output", str(output_path),
            "--project-file", str(project_file),
        ]
    )

    assert rc == 0
    html = output_path.read_text(encoding="utf-8")
    assert "low relevance to your research question" in html


def _citing_index(index_path, refs):
    entry = LibraryEntry(
        file_path="a.pdf",
        file_hash="hash-a",
        file_type="pdf",
        size_bytes=100,
        indexed_at="2026-01-01T00:00:00",
        confidence="verified-doi",
        doi="10.1/a",
        paper=Paper(title="Sleep deprivation in adolescents", doi="10.1/a", year=2023, sources=["crossref"]),
        references=refs,
    )
    LibraryIndex([entry]).save(index_path)


def test_visualize_library_filters_off_topic_citations_by_default(tmp_path):
    index_path = tmp_path / "library" / "index.json"
    output_path = tmp_path / "library" / "visualization.html"
    _citing_index(index_path, [
        {"doi": "10.9/coffee", "title": "Coffee bean price volatility in Brazil", "year": 2018},
        {"doi": "10.9/sleep", "title": "Adolescent sleep and school schedules", "year": 2005},
    ])

    rc = main([
        "visualize-library",
        "--index-path", str(index_path),
        "--output", str(output_path),
        "--project-file", str(tmp_path / "project.json"),
        "--question", "Does sleep deprivation affect adolescent decision-making?",
    ])

    assert rc == 0
    html = output_path.read_text(encoding="utf-8")
    assert "Adolescent sleep and school schedules" in html
    assert "Coffee bean price volatility in Brazil" not in html


def test_visualize_library_min_gap_relevance_zero_keeps_everything(tmp_path):
    index_path = tmp_path / "library" / "index.json"
    output_path = tmp_path / "library" / "visualization.html"
    _citing_index(index_path, [
        {"doi": "10.9/coffee", "title": "Coffee bean price volatility in Brazil", "year": 2018},
    ])

    rc = main([
        "visualize-library",
        "--index-path", str(index_path),
        "--output", str(output_path),
        "--project-file", str(tmp_path / "project.json"),
        "--question", "Does sleep deprivation affect adolescent decision-making?",
        "--min-gap-relevance", "0",
    ])

    assert rc == 0
    assert "Coffee bean price volatility in Brazil" in output_path.read_text(encoding="utf-8")


def _one_paper_index(index_path):
    index = LibraryIndex(
        [
            LibraryEntry(
                file_path="a.pdf",
                file_hash="hash-a",
                file_type="pdf",
                size_bytes=100,
                indexed_at="2026-01-01T00:00:00",
                confidence="verified-doi",
                doi="10.1/a",
                paper=Paper(title="A Paper", doi="10.1/a", year=2023, sources=["crossref"]),
            )
        ]
    )
    index.save(index_path)
    return index_path


def test_the_visualization_is_a_living_view_not_a_versioned_one(tmp_path):
    """It describes the library as it stands now. An old snapshot of a
    library you have since changed is not history — it is a stale picture of
    something that no longer exists."""
    index_path = _one_paper_index(tmp_path / "library" / "index.json")
    output_path = tmp_path / "library" / "visualization.html"

    for _ in range(3):
        assert main(["visualize-library", "--index-path", str(index_path), "--output", str(output_path)]) == 0

    assert output_path.is_file()
    assert not (tmp_path / "library" / "previous").exists()
    assert sorted(p.name for p in (tmp_path / "library").iterdir() if p.is_file()) == [
        "index.json",
        "literature_matrix.xlsx",
        "visualization.html",
    ]
